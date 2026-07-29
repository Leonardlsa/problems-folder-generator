from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import openpyxl
import pytest
from PIL import Image
from docx import Document

from core import GenerateError, generate


class _Log:
    def __init__(self) -> None:
        self.lines: list[str] = []

    def append(self, text: str) -> None:
        self.lines.append(text)


class _Widget:
    def __init__(self, confirm: bool = False) -> None:
        self.textBrowser = _Log()
        self.confirm = confirm
        self.confirmations: list[tuple[str, str, list[int]]] = []

    def confirm_missing_images(self, name: str, version_subject: str, missing: list[int]) -> bool:
        self.confirmations.append((name, version_subject, missing))
        return self.confirm


def _image(path: Path, rgb: tuple[int, int, int]) -> None:
    Image.new("RGB", (120, 80), rgb).save(path)


def _feedback(path: Path, rows: list[tuple[str, str]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["姓名", "错题号"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _media_hashes(docx_path: Path) -> set[str]:
    hashes: set[str] = set()
    with zipfile.ZipFile(docx_path) as archive:
        for name in archive.namelist():
            if name.startswith("word/media/"):
                hashes.add(hashlib.sha256(archive.read(name)).hexdigest())
    return hashes


def _file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_generate_coupled_student_and_teacher_versions(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    student_images = tmp_path / "student-images"
    teacher_images = tmp_path / "teacher-images"
    student_images.mkdir()
    teacher_images.mkdir()

    _image(student_images / "01.png", (255, 0, 0))
    _image(student_images / "02.png", (0, 255, 0))
    _image(teacher_images / "01.png", (0, 0, 255))
    _image(teacher_images / "02.png", (255, 255, 0))
    feedback = tmp_path / "feedback.xlsx"
    _feedback(feedback, [("张三", "1"), ("李四", "2")])

    widget = _Widget()
    generate(
        str(student_images),
        str(teacher_images),
        str(feedback),
        "高一一班",
        "物理",
        "03",
        "0729",
        widget,
    )

    zhang_student = tmp_path / "学生/张三/物理学生版/第03周错题集.docx"
    zhang_teacher = tmp_path / "学生/张三/物理教师版/第03周错题集.docx"
    li_student = tmp_path / "学生/李四/物理学生版/第03周错题集.docx"
    li_teacher = tmp_path / "学生/李四/物理教师版/第03周错题集.docx"
    for path in (zhang_student, zhang_teacher, li_student, li_teacher):
        assert path.is_file()

    assert _media_hashes(zhang_student) == {_file_hash(student_images / "01.png")}
    assert _media_hashes(zhang_teacher) == {_file_hash(teacher_images / "01.png")}
    assert _media_hashes(li_student) == {_file_hash(student_images / "02.png")}
    assert _media_hashes(li_teacher) == {_file_hash(teacher_images / "02.png")}

    assert "物理学生版" in Document(zhang_student).sections[0].header.paragraphs[0].text
    assert "物理教师版" in Document(zhang_teacher).sections[0].header.paragraphs[0].text


def test_all_correct_and_unsubmitted_are_mirrored(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    student_images = tmp_path / "s"
    teacher_images = tmp_path / "t"
    student_images.mkdir()
    teacher_images.mkdir()
    for number in (1, 2):
        _image(student_images / f"{number:02d}.png", (number * 40, 0, 0))
        _image(teacher_images / f"{number:02d}.png", (0, number * 40, 0))

    feedback = tmp_path / "feedback.xlsx"
    _feedback(feedback, [("甲", "全对"), ("乙", "未交")])
    generate(str(student_images), str(teacher_images), str(feedback), "一班", "数学", "1", "0729", _Widget())

    for suffix in ("学生版", "教师版"):
        all_correct = Document(tmp_path / f"学生/甲/数学{suffix}/第1周错题集.docx")
        assert any(paragraph.text == "全对" for paragraph in all_correct.paragraphs)
        unsubmitted = tmp_path / f"学生/乙/数学{suffix}/第1周错题集.docx"
        assert len(_media_hashes(unsubmitted)) == 2


def test_image_sets_must_have_matching_problem_numbers(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    student_images = tmp_path / "s"
    teacher_images = tmp_path / "t"
    student_images.mkdir()
    teacher_images.mkdir()
    _image(student_images / "01.png", (1, 2, 3))
    _image(student_images / "02.png", (4, 5, 6))
    _image(teacher_images / "01.png", (7, 8, 9))
    feedback = tmp_path / "feedback.xlsx"
    _feedback(feedback, [("甲", "1")])

    with pytest.raises(GenerateError, match="图片题号不对应"):
        generate(str(student_images), str(teacher_images), str(feedback), "一班", "物理", "1", "0729", _Widget())


def test_missing_requested_problem_can_be_cancelled_before_saving(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    student_images = tmp_path / "s"
    teacher_images = tmp_path / "t"
    student_images.mkdir()
    teacher_images.mkdir()
    _image(student_images / "01.png", (1, 2, 3))
    _image(teacher_images / "01.png", (4, 5, 6))
    feedback = tmp_path / "feedback.xlsx"
    _feedback(feedback, [("甲", "2")])

    with pytest.raises(GenerateError, match="缺少图片 2"):
        generate(str(student_images), str(teacher_images), str(feedback), "一班", "物理", "1", "0729", _Widget())

    assert not list((tmp_path / "学生").rglob("*.docx"))


def test_existing_week_document_is_appended_and_previous_tree_is_backed_up(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    student_images = tmp_path / "s"
    teacher_images = tmp_path / "t"
    student_images.mkdir()
    teacher_images.mkdir()
    _image(student_images / "01.png", (1, 2, 3))
    _image(teacher_images / "01.png", (4, 5, 6))
    feedback = tmp_path / "feedback.xlsx"
    _feedback(feedback, [("甲", "全对")])

    args = (str(student_images), str(teacher_images), str(feedback), "一班", "物理", "1")
    generate(*args, "0728", _Widget())
    generate(*args, "0729", _Widget())

    path = tmp_path / "学生/甲/物理学生版/第1周错题集.docx"
    headings = [paragraph.text for paragraph in Document(path).paragraphs if paragraph.style.name.startswith("Heading")]
    assert headings == ["0728", "0729"]
    assert (tmp_path / "backup/甲/物理学生版/第1周错题集.docx").is_file()
    assert (tmp_path / "backup/甲/物理教师版/第1周错题集.docx").is_file()
